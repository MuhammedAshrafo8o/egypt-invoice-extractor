"""End-to-end tests: real HTTP server, real database, simulated WhatsApp and ETA.

Run with:  python -m unittest discover tests
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from datetime import date, timedelta
from pathlib import Path

import requests
import uvicorn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.config import load_settings  # noqa: E402
from app.main import create_app  # noqa: E402
from invoice_extractor.eta_client import map_document  # noqa: E402
from invoice_extractor.extractors.pdf_text import fix_reversed_arabic, looks_reversed  # noqa: E402

SAMPLES = ROOT / "samples" / "invoices"
TODAY = date.today()
PERIOD = TODAY.strftime("%Y-%m")
APP_SECRET = "test-app-secret"
VERIFY_TOKEN = "verify-me"


# --------------------------------------------------------------------- fakes
class FakeWhatsApp:
    def __init__(self):
        self.sent: list[tuple[str, str]] = []
        self.media: dict[str, tuple[bytes, str]] = {}

    def download_media(self, media_id):
        return self.media[media_id]

    def send_text(self, to, text):
        self.sent.append((to, text))


def _raw_doc(uuid, issuer_id, receiver_id, number, net, vat, issued):
    document = {
        "issuer": {"id": issuer_id, "name": f"Issuer {issuer_id}"},
        "receiver": {"id": receiver_id, "name": f"Receiver {receiver_id}"},
        "internalID": number, "dateTimeIssued": issued,
        "taxTotals": [{"taxType": "T1", "amount": vat}],
        "netAmount": net, "totalAmount": net + vat,
        "invoiceLines": [{"unitValue": {"currencySold": "EGP"}}],
    }
    return {"uuid": uuid, "status": "Valid", "document": json.dumps(document)}


class FakeETA:
    calls: list = []
    issued = f"{(TODAY - timedelta(days=2)).isoformat()}T10:00:00Z"
    docs = {
        "UUIDSALE1": _raw_doc("UUIDSALE1", "300112455", "777888999", "S-1", 1000.0, 140.0, issued),
        "UUIDBUY1": _raw_doc("UUIDBUY1", "512334901", "300112455", "P-9", 500.0, 70.0, issued),
        "SAMPLE7Q2K9XW4D1M8FZ0E3R6T5Y": _raw_doc("SAMPLE7Q2K9XW4D1M8FZ0E3R6T5Y", "205778341", "300112455",
                                                 "EI-55021", 7250.0, 1015.0, issued),
    }

    def __init__(self, client_id, client_secret, env="preprod"):
        FakeETA.calls.append((client_id, env))

    def search_documents(self, start, end, direction=None, page_size=100):
        yield {"uuid": "UUIDSALE1", "status": "Valid"}
        yield {"uuid": "UUIDBUY1", "status": "Valid"}
        yield {"uuid": "CANCELLED", "status": "Cancelled"}

    def get_document(self, uuid):
        return self.docs[uuid]

    def fetch_fields(self, uuid):
        return map_document(self.get_document(uuid))


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class AppTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        subprocess.run([sys.executable, str(ROOT / "samples" / "make_samples.py")], check=True, capture_output=True)
        cls.tmp = Path(tempfile.mkdtemp())
        import os
        os.environ.update(DATA_DIR=str(cls.tmp), USE_AI="false", USE_OCR="false", WHATSAPP_APP_SECRET=APP_SECRET,
                          WHATSAPP_VERIFY_TOKEN=VERIFY_TOKEN, ALLOW_SIGNUP="true")
        os.environ.pop("ANTHROPIC_API_KEY", None)
        cls.settings = load_settings()
        cls.wa = FakeWhatsApp()
        cls.app = create_app(cls.settings, eta_client_factory=FakeETA, whatsapp_client=cls.wa)
        cls.db = cls.app.state.db
        cls.port = _free_port()
        cls.base = f"http://127.0.0.1:{cls.port}"
        config = uvicorn.Config(cls.app, host="127.0.0.1", port=cls.port, log_level="warning")
        cls.server = uvicorn.Server(config)
        cls.thread = threading.Thread(target=cls.server.run, daemon=True)
        cls.thread.start()
        for _ in range(100):
            try:
                requests.get(cls.base + "/health", timeout=1)
                break
            except requests.ConnectionError:
                time.sleep(0.05)

    @classmethod
    def tearDownClass(cls):
        cls.server.should_exit = True
        cls.thread.join(timeout=5)
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ---------------------------------------------------------------- helpers
    def session(self, email, name="مكتب", password="password123"):
        s = requests.Session()
        r = s.post(self.base + "/register", data={"name": name, "email": email, "password": password})
        if "متسجل قبل كده" in r.text:
            r = s.post(self.base + "/login", data={"email": email, "password": password})
        self.assertEqual(r.status_code, 200)
        return s

    @staticmethod
    def csrf(s, url):
        html = s.get(url).text
        return re.search(r'name="csrf" value="([^"]+)"', html).group(1)

    def create_client(self, s, name, tax_id="", phone=""):
        token = self.csrf(s, self.base + "/")
        r = s.post(self.base + "/clients", data={"csrf": token, "name": name, "tax_id": tax_id,
                                                 "whatsapp_phone": phone}, allow_redirects=False)
        self.assertEqual(r.status_code, 303, r.text[:300])
        return int(r.headers["location"].rsplit("/", 1)[1])

    def upload(self, s, client_id, *names):
        token = self.csrf(s, f"{self.base}/clients/{client_id}")
        files = [("files", (n, (SAMPLES / n).read_bytes(), "application/octet-stream")) for n in names]
        r = s.post(f"{self.base}/clients/{client_id}/upload", data={"csrf": token, "period": PERIOD}, files=files)
        self.assertEqual(r.status_code, 200)
        self.wait_for_reading(client_id)
        return r

    def wait_for_reading(self, client_id, timeout=60):
        """Several files are read in the background; wait until they're all done."""
        deadline = time.time() + timeout
        while self.db.pending_jobs(client_id) and time.time() < deadline:
            time.sleep(0.1)
        self.assertEqual(self.db.pending_jobs(client_id), 0, "background reading did not finish")

    def invoices(self, client_id):
        return {i["original_filename"]: i for i in self.db.list_invoices(client_id)}


class AuthAndAccessTests(AppTestCase):
    def test_login_required_and_wrong_password(self):
        r = requests.get(self.base + "/", allow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.assertIn("/login", r.headers["location"])
        self.session("a@x.com")
        r = requests.post(self.base + "/login", data={"email": "a@x.com", "password": "wrong-pass"})
        self.assertEqual(r.status_code, 400)

    def test_accountants_cannot_see_each_other(self):
        s1 = self.session("owner@x.com")
        cid = self.create_client(s1, "عميل خاص")
        self.upload(s1, cid, "01_digital_ok.pdf")
        inv_id = self.db.list_invoices(cid)[0]["id"]
        s2 = self.session("other@x.com")
        for url in (f"/clients/{cid}", f"/invoices/{inv_id}", f"/invoices/{inv_id}/file",
                    f"/clients/{cid}/export.xlsx"):
            self.assertEqual(s2.get(self.base + url).status_code, 404, url)
        self.assertNotIn("عميل خاص", s2.get(self.base + "/").text)

    def test_register_errors_are_shown_not_crashing(self):
        self.session("dup@x.com", name="مكتب")
        r = requests.post(self.base + "/register", data={"name": "مكتب تاني", "email": "dup@x.com",
                                                         "password": "password123"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("الإيميل ده متسجل قبل كده", r.text)
        r = requests.post(self.base + "/register", data={"name": "", "email": "bad", "password": "1"})
        self.assertEqual(r.status_code, 400)

    def test_csrf_required(self):
        s = self.session("csrf@x.com")
        r = s.post(self.base + "/clients", data={"name": "بدون توكن"})
        self.assertIn("انتهت صلاحية الصفحة", r.text)
        self.assertNotIn("بدون توكن", s.get(self.base + "/").text)


class InvoiceFlowTests(AppTestCase):
    def test_upload_review_fix_approve_export(self):
        s = self.session("flow@x.com")
        cid = self.create_client(s, "شركة الدلتا للتجارة", "300112455")
        r = self.upload(s, cid, "01_digital_ok.pdf", "02_total_mismatch.pdf", "07_arabic.pdf",
                        "08_arabic_visual_order.pdf", "05_photo.jpg")
        self.assertIn("اترفع 5 ملف", r.text)
        inv = self.invoices(cid)
        self.assertEqual(inv["01_digital_ok.pdf"]["status"], "ok")
        self.assertEqual(inv["01_digital_ok.pdf"]["direction"], "purchase")  # client is the buyer
        self.assertEqual(inv["02_total_mismatch.pdf"]["status"], "error")
        self.assertEqual(inv["07_arabic.pdf"]["status"], "ok")
        self.assertEqual(inv["07_arabic.pdf"]["supplier_name"], "شركة الأهرام للتوريدات الهندسية")
        dup = inv["08_arabic_visual_order.pdf"]
        self.assertEqual(dup["status"], "error")
        self.assertIn("duplicate", [i["code"] for i in dup["issues"]])
        self.assertEqual(inv["05_photo.jpg"]["status"], "error")  # AI off → nothing read

        # Review page renders, PDF page image renders.
        bad = inv["02_total_mismatch.pdf"]
        page = s.get(f"{self.base}/invoices/{bad['id']}")
        self.assertIn("الإجمالي 3,050.00", page.text)
        img = s.get(f"{self.base}/invoices/{bad['id']}/page/0.png")
        self.assertEqual(img.headers["content-type"], "image/png")
        self.assertTrue(img.content.startswith(b"\x89PNG"))

        # Approving an invoice with errors is refused.
        token = self.csrf(s, f"{self.base}/invoices/{bad['id']}")
        s.post(f"{self.base}/invoices/{bad['id']}/edit", data={"csrf": token, "action": "approve"})
        self.assertEqual(self.db.get_invoice(bad["id"])["approved"], 0)

        # Fix the total, save and approve.
        form = {k: bad[k] or "" for k in ("supplier_name", "supplier_tax_id", "buyer_name", "buyer_tax_id",
                                          "invoice_number", "invoice_date", "currency", "subtotal", "vat_amount")}
        form.update(csrf=token, total="2850.00", direction="purchase", period=PERIOD, action="save_approve")
        r = s.post(f"{self.base}/invoices/{bad['id']}/edit", data=form)
        self.assertIn("واتعتمدت", r.text)
        fixed = self.db.get_invoice(bad["id"])
        self.assertEqual((fixed["status"], fixed["approved"], fixed["total"]), ("ok", 1, 2850.0))

        # Delete the duplicate.
        token = self.csrf(s, f"{self.base}/invoices/{dup['id']}")
        s.post(f"{self.base}/invoices/{dup['id']}/edit", data={"csrf": token, "action": "delete"})
        self.assertIsNone(self.db.get_invoice(dup["id"]))

        # Summary: input VAT = 1400 + 350 + 672 (photo has no numbers and stays an error).
        summary = self.app.state.service.summary(cid, PERIOD)
        self.assertEqual(summary["purchases"]["vat"], 2422.0)
        self.assertEqual(summary["net_vat"], -2422.0)

        # Excel export.
        from openpyxl import load_workbook
        r = s.get(f"{self.base}/clients/{cid}/export.xlsx?period={PERIOD}")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["ملخص الإقرار", "المشتريات", "المبيعات", "المشتريات حسب المورد"])
        purchases = wb["المشتريات"]
        self.assertEqual(purchases.max_row, 1 + 4 + 1)  # header + 4 invoices + totals
        self.assertTrue(str(purchases.cell(row=6, column=8).value).startswith("=SUMIFS"))
        self.assertEqual(wb["ملخص الإقرار"]["B11"].value, "=B8-B10")
        # The "counted" column must give exactly the system's totals (what SUMIFS will add up).
        counted_vat = sum(r[7] or 0 for r in purchases.iter_rows(min_row=2, max_row=5, values_only=True) if r[11] == "نعم")
        self.assertAlmostEqual(counted_vat, summary["purchases"]["vat"])

    def test_single_upload_opens_the_invoice(self):
        s = self.session("single@x.com")
        cid = self.create_client(s, "عميل", "300112455")
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        r = s.post(f"{self.base}/clients/{cid}/upload", data={"csrf": token, "period": PERIOD},
                   files=[("files", ("07_arabic.pdf", (SAMPLES / "07_arabic.pdf").read_bytes(), "application/pdf"))],
                   allow_redirects=False)
        self.assertEqual(r.status_code, 303)
        self.assertRegex(r.headers["location"], r"^/invoices/\d+$")
        page = s.get(self.base + r.headers["location"]).text
        self.assertIn("اتقرت الفاتورة", page)
        self.assertIn('value="488201736"', page)  # fields are on screen

    def test_period_for_suspicious_dates(self):
        from app.services import current_period, period_for
        self.assertEqual(period_for(TODAY.isoformat()), PERIOD)
        self.assertEqual(period_for("2062-10-12"), current_period())   # misread year
        self.assertEqual(period_for("2019-01-05"), current_period())   # far too old
        self.assertEqual(period_for(None), current_period())
        last_month = (TODAY.replace(day=1) - timedelta(days=1))
        self.assertEqual(period_for(last_month.isoformat()), last_month.strftime("%Y-%m"))

    def test_sale_detection_and_filters(self):
        s = self.session("seller@x.com")
        cid = self.create_client(s, "مكتبة النيل", "512334901")  # the issuer of sample 01
        self.upload(s, cid, "01_digital_ok.pdf")
        inv = self.db.list_invoices(cid)[0]
        self.assertEqual(inv["direction"], "sale")
        html = s.get(f"{self.base}/clients/{cid}?period={PERIOD}&direction=purchase").text
        self.assertIn("مفيش فواتير هنا", html)

    def test_rejects_unsupported_file(self):
        s = self.session("files@x.com")
        cid = self.create_client(s, "عميل")
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        r = s.post(f"{self.base}/clients/{cid}/upload", data={"csrf": token},
                   files=[("files", ("notes.txt", b"hello", "text/plain"))])
        self.assertIn("مش مدعوم", r.text)
        self.assertEqual(self.db.list_invoices(cid), [])


class ExcelExportTests(unittest.TestCase):
    """The counted flag in the workbook must follow the same rule as the on-screen summary."""

    def test_counted_rule_matches_summary_on_edge_cases(self):
        from app.exports import build_workbook, is_counted
        from openpyxl import load_workbook

        def inv(i, **kw):
            base = dict(id=i, status="ok", approved=0, direction="purchase", currency="EGP", issues=[], channel="upload",
                        invoice_date="2026-09-10", invoice_number=f"N{i}", supplier_name=f"S{i}", supplier_tax_id=f"{i}" * 9,
                        buyer_name=None, buyer_tax_id=None, subtotal=100.0 * i, vat_rate=0.14, vat_amount=14.0 * i,
                        total=114.0 * i)
            base.update(kw)
            return base

        rows = [inv(1), inv(2, status="review"), inv(3, status="error", approved=1), inv(4, status="error"),
                inv(5, currency="USD"), inv(6, direction="sale", buyer_name="B", buyer_tax_id="999999999")]
        expected_in = [1, 2, 3, 6]
        self.assertEqual([r["id"] for r in rows if is_counted(r)], expected_in)
        data = build_workbook({"name": "X", "tax_id": "300112455"}, "2026-09", rows, {"review_count": 1})
        wb = load_workbook(io.BytesIO(data))
        flags = {r[2]: r[11] for r in wb["المشتريات"].iter_rows(min_row=2, values_only=True) if r[2]}
        self.assertEqual(flags, {"N1": "نعم", "N2": "نعم", "N3": "نعم", "N4": "لا", "N5": "لا"})
        self.assertEqual(wb["المشتريات"]["B2"].number_format, "dd/mm/yyyy")
        self.assertTrue(wb["المشتريات"].auto_filter.ref)


class AccountantValueTests(AppTestCase):
    """Features that save the accountant time beyond reading invoices."""

    def test_party_memory_fills_names_and_flags_near_miss_tax_ids(self):
        from invoice_extractor.models import Invoice
        s = self.session("memory@x.com")
        cid = self.create_client(s, "شركة الدلتا", "300112455")
        client = self.db.get_client(cid)
        self.upload(s, cid, "01_digital_ok.pdf")
        inv = self.db.list_invoices(cid)[0]
        self.app.state.service.approve(inv)  # supplier 512334901 is now known
        service = self.app.state.service

        known = Invoice(source_file="x", supplier_tax_id="512334901")
        service._apply_party_memory(client, known)
        self.assertEqual(known.supplier_name, "Nile Office Supplies LLC")

        typo = Invoice(source_file="y", supplier_tax_id="512334907", supplier_name="Nile Office")
        service._apply_party_memory(client, typo)
        self.assertIn("supplier_tax_id_near_known", [i.code for i in typo.issues])
        self.assertEqual(typo.status, "review")

        unrelated = Invoice(source_file="z", supplier_tax_id="777888999")
        service._apply_party_memory(client, unrelated)
        self.assertEqual(unrelated.issues, [])

    def test_bulk_approve_skips_errors(self):
        s = self.session("bulk@x.com")
        cid = self.create_client(s, "عميل", "300112455")
        self.upload(s, cid, "04_eta_einvoice.pdf", "02_total_mismatch.pdf")
        inv = {i["original_filename"]: i for i in self.db.list_invoices(cid)}
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        r = s.post(f"{self.base}/clients/{cid}/approve", data={
            "csrf": token, "period": PERIOD,
            "ids": [inv["04_eta_einvoice.pdf"]["id"], inv["02_total_mismatch.pdf"]["id"]]})
        self.assertIn("اتعتمدت 1 فاتورة", r.text)
        self.assertIn("ماتعتمدتش", r.text)
        self.assertEqual(self.db.get_invoice(inv["04_eta_einvoice.pdf"]["id"])["approved"], 1)
        self.assertEqual(self.db.get_invoice(inv["02_total_mismatch.pdf"]["id"])["approved"], 0)

    def test_client_report_and_whatsapp_links(self):
        s = self.session("report@x.com", name="مكتب الأمل")
        cid = self.create_client(s, "شركة الدلتا", "300112455", "01022223333")
        quiet = self.create_client(s, "عميل ساكت", "", "01044445555")
        self.upload(s, cid, "01_digital_ok.pdf", "02_total_mismatch.pdf")

        report = s.get(f"{self.base}/clients/{cid}/report?period={PERIOD}").text
        self.assertIn("مكتب الأمل", report)
        self.assertIn("رصيد دائن", report)            # purchases only → credit
        self.assertIn("محتاجين منك (1)", report)       # only the invoice with a wrong total
        self.assertNotIn("القراءة المجانية", report)    # internal notes never reach the client
        self.assertIn("1,400.00", report)

        home = s.get(self.base + "/").text
        self.assertIn("لسه مابعتوش فواتير", home)
        self.assertIn("https://wa.me/201044445555?text=", home)
        self.assertIn("وقت وفّره السيستم", home)

        bad = next(i for i in self.db.list_invoices(cid) if i["status"] == "error")
        page = s.get(f"{self.base}/invoices/{bad['id']}").text
        self.assertIn("https://wa.me/201022223333?text=", page)
        self.assertIn("اسأل العميل على واتساب", page)

    def test_review_queue_steps_to_next_invoice(self):
        s = self.session("queue@x.com")
        cid = self.create_client(s, "عميل", "300112455")
        self.upload(s, cid, "04_eta_einvoice.pdf", "09_arabic_photo.jpg", "07_arabic.pdf")
        queue = s.get(self.base + "/review").text
        ids = list(dict.fromkeys(re.findall(r"/invoices/(\d+)\?queue=1", queue)))
        self.assertGreaterEqual(len(ids), 1)
        page = s.get(f"{self.base}/invoices/{ids[0]}?queue=1").text
        self.assertIn("فاتورة 1 من", page)


class SearchAndAnalyticsTests(AppTestCase):
    def setUp(self):
        # A fresh account per test so invoices don't pile up across tests.
        self.s = self.session(f"search-{self._testMethodName}@x.com")
        self.cid = self.create_client(self.s, "شركة الدلتا", "300112455")
        self.upload(self.s, self.cid, "07_arabic.pdf", "01_digital_ok.pdf", "02_total_mismatch.pdf")

    def test_arabic_insensitive_search_and_filters(self):
        page = self.s.get(self.base + "/invoices", params={"q": "الاهرام"}).text  # stored as «الأهرام»
        self.assertIn("1 فاتورة", page)
        self.assertIn("شركة الأهرام", page)
        for q in ("512-334-901", "١١٤٠٠", "11,400", "inv-2026"):
            self.assertIn("1 فاتورة", self.s.get(self.base + "/invoices", params={"q": q}).text, q)
        errors = self.s.get(self.base + "/invoices", params={"status": "error"}).text
        self.assertIn("1 فاتورة", errors)
        self.assertIn("Cairo Print House", errors)
        big = self.s.get(self.base + "/invoices", params={"min": "6000"}).text
        self.assertIn("1 فاتورة", big)  # only the 11,400 invoice
        other = self.session("someone-else@x.com")
        self.assertIn("0 فاتورة", other.get(self.base + "/invoices", params={"q": "الاهرام"}).text)

    def test_csv_export_and_quick_search(self):
        r = self.s.get(self.base + "/invoices/export.csv", params={"q": "الأهرام"})
        self.assertTrue(r.content.startswith(b"\xef\xbb\xbf"))  # BOM so Excel shows Arabic
        text = r.content.decode("utf-8-sig")
        self.assertEqual(len(text.strip().splitlines()), 2)
        self.assertIn("488201736", text)
        data = self.s.get(self.base + "/api/search", params={"q": "دلتا"}).json()
        self.assertEqual(data["clients"][0]["name"], "شركة الدلتا")
        data = self.s.get(self.base + "/api/search", params={"q": "Nile"}).json()
        self.assertTrue(data["invoices"][0]["title"].startswith("Nile Office Supplies LLC"))

    def test_analytics_page(self):
        page = self.s.get(self.base + "/analytics").text
        self.assertIn("ضريبة المخرجات والمدخلات", page)
        self.assertIn("<svg", page)
        self.assertIn("أكبر الموردين في ضريبة المدخلات", page)
        self.assertIn("الإجمالي مش مظبوط", page)  # problem category from the wrong-total invoice
        self.assertIn("Nile Office Supplies LLC", page)


class WhatsAppTests(AppTestCase):
    def post_webhook(self, payload, secret=APP_SECRET):
        body = json.dumps(payload).encode()
        sig = "sha256=" + hmac.new(secret.encode(), body, hashlib.sha256).hexdigest()
        return requests.post(self.base + "/webhooks/whatsapp", data=body,
                             headers={"Content-Type": "application/json", "X-Hub-Signature-256": sig})

    @staticmethod
    def message(msg_id, sender, kind="document", media_id="m1", filename="inv.pdf", mime="application/pdf"):
        msg = {"id": msg_id, "from": sender, "type": kind}
        if kind in ("document", "image"):
            msg[kind] = {"id": media_id, "mime_type": mime, **({"filename": filename} if kind == "document" else {})}
        if kind == "text":
            msg["text"] = {"body": "السلام عليكم"}
        return {"object": "whatsapp_business_account",
                "entry": [{"changes": [{"value": {"messages": [msg]}}]}]}

    def wait_for(self, predicate, timeout=10):
        end = time.time() + timeout
        while time.time() < end:
            if predicate():
                return True
            time.sleep(0.05)
        return False

    def test_verify_handshake(self):
        ok = requests.get(self.base + "/webhooks/whatsapp", params={
            "hub.mode": "subscribe", "hub.verify_token": VERIFY_TOKEN, "hub.challenge": "12345"})
        self.assertEqual(ok.text, "12345")
        bad = requests.get(self.base + "/webhooks/whatsapp", params={
            "hub.mode": "subscribe", "hub.verify_token": "nope", "hub.challenge": "1"})
        self.assertEqual(bad.status_code, 403)

    def test_bad_signature_rejected(self):
        r = self.post_webhook(self.message("x1", "201000000000"), secret="wrong")
        self.assertEqual(r.status_code, 401)

    def test_invoice_from_registered_client(self):
        s = self.session("wa@x.com")
        cid = self.create_client(s, "عميل واتساب", "300112455", "01011112222")
        self.wa.media["media-42"] = ((SAMPLES / "07_arabic.pdf").read_bytes(), "application/pdf")
        payload = self.message("wamid.A1", "201011112222", media_id="media-42", filename="فاتورة.pdf")
        self.assertEqual(self.post_webhook(payload).status_code, 200)
        self.assertTrue(self.wait_for(lambda: len(self.db.list_invoices(cid)) == 1))
        inv = self.db.list_invoices(cid)[0]
        self.assertEqual((inv["channel"], inv["status"], inv["total"]), ("whatsapp", "ok", 5472.0))
        self.assertTrue(self.wait_for(lambda: any(to == "201011112222" and "5,472.00" in t
                                                  for to, t in self.wa.sent)))
        # Meta retries the same message → ignored.
        self.post_webhook(payload)
        time.sleep(0.5)
        self.assertEqual(len(self.db.list_invoices(cid)), 1)

    def test_unknown_sender_and_text(self):
        self.post_webhook(self.message("wamid.U1", "201099999999"))
        self.assertTrue(self.wait_for(lambda: any(to == "201099999999" and "مش مسجل" in t
                                                  for to, t in self.wa.sent)))
        s = self.session("wa2@x.com")
        self.create_client(s, "عميل نص", "", "01033334444")
        self.post_webhook(self.message("wamid.T1", "201033334444", kind="text"))
        self.assertTrue(self.wait_for(lambda: any(to == "201033334444" and "ابعت صورة" in t
                                                  for to, t in self.wa.sent)))


class ETATests(AppTestCase):
    def connect_eta(self, s, cid):
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        s.post(f"{self.base}/clients/{cid}/edit", data={
            "csrf": token, "name": "شركة الدلتا", "tax_id": "300112455", "eta_client_id": "cid",
            "eta_client_secret": "secret"})

    def test_sync_month(self):
        s = self.session("eta@x.com")
        cid = self.create_client(s, "شركة الدلتا", "300112455")
        self.connect_eta(s, cid)
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        r = s.post(f"{self.base}/clients/{cid}/eta-sync", data={"csrf": token, "period": PERIOD})
        self.assertIn("اتسحب 2 فاتورة", r.text)
        inv = {i["eta_uuid"]: i for i in self.db.list_invoices(cid)}
        self.assertEqual(inv["UUIDSALE1"]["direction"], "sale")
        self.assertEqual(inv["UUIDBUY1"]["direction"], "purchase")
        self.assertTrue(all(i["approved"] and i["channel"] == "eta_sync" for i in inv.values()))
        summary = self.app.state.service.summary(cid, PERIOD)
        self.assertEqual((summary["sales"]["vat"], summary["purchases"]["vat"], summary["net_vat"]),
                         (140.0, 70.0, 70.0))
        # Second sync imports nothing new.
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        r = s.post(f"{self.base}/clients/{cid}/eta-sync", data={"csrf": token, "period": PERIOD})
        self.assertIn("اتسحب 0 فاتورة", r.text)
        # The secret is never echoed back into the page.
        self.assertNotIn('value="secret"', s.get(f"{self.base}/clients/{cid}").text)

    def test_qr_invoice_pulled_from_eta(self):
        s = self.session("etaqr@x.com")
        cid = self.create_client(s, "شركة الدلتا", "300112455")
        self.connect_eta(s, cid)
        self.upload(s, cid, "04_eta_einvoice.pdf")
        inv = self.db.list_invoices(cid)[0]
        self.assertEqual((inv["source_type"], inv["status"], inv["vat_amount"]), ("eta_api", "ok", 1015.0))


class BulkUploadTests(AppTestCase):
    def post_files(self, s, cid, files):
        token = self.csrf(s, f"{self.base}/clients/{cid}")
        return s.post(f"{self.base}/clients/{cid}/upload", data={"csrf": token, "period": PERIOD}, files=[
            ("files", (name, data, "application/octet-stream")) for name, data in files])

    def test_many_files_each_one_invoice_and_no_double_upload(self):
        s = self.session("bulk@x.com")
        cid = self.create_client(s, "شركة الدلتا", "300112455")
        names = ["01_digital_ok.pdf", "02_total_mismatch.pdf", "07_arabic.pdf", "09_arabic_photo.jpg",
                 "05_photo.jpg", "06_scanned.pdf"]
        files = [(n, (SAMPLES / n).read_bytes()) for n in names]
        files.append(("copy_of_01.pdf", files[0][1]))                    # same bytes, other name
        r = self.post_files(s, cid, files)
        self.assertIn("اترفع 6 ملف", r.text)
        self.assertIn("متكرر في نفس المجموعة", r.text)
        self.wait_for_reading(cid)
        self.assertEqual(sorted(self.invoices(cid)), sorted(names))   # one invoice per file
        progress = s.get(f"{self.base}/clients/{cid}/progress.json").json()
        self.assertEqual((progress["total"], progress["done"], progress["running"]), (6, 6, False))
        # Uploading the same files again adds nothing.
        r = self.post_files(s, cid, files[:3])
        self.assertIn("اترفع قبل كده", r.text)
        self.wait_for_reading(cid)
        self.assertEqual(len(self.db.list_invoices(cid)), 6)
        # Same for a single file (read right away).
        r = self.post_files(s, cid, files[:1])
        self.assertIn("اترفع قبل كده", r.text)
        self.assertEqual(len(self.db.list_invoices(cid)), 6)

    def test_upload_limit(self):
        s = self.session("bulk-limit@x.com")
        cid = self.create_client(s, "شركة الدلتا", "300112455")
        settings = self.app.state.settings
        old, settings.max_files_per_upload = settings.max_files_per_upload, 2
        try:
            r = self.post_files(s, cid, [(f"{n}.pdf", (SAMPLES / "01_digital_ok.pdf").read_bytes() + bytes([n]))
                                         for n in range(3)])
        finally:
            settings.max_files_per_upload = old
        self.assertIn("والحد 2 ملف في المرة", r.text)
        self.assertEqual(self.db.list_invoices(cid), [])

    def test_invoice_not_in_client_name_is_flagged(self):
        s = self.session("bulk-other@x.com")
        cid = self.create_client(s, "شركة تانية", "999888777")
        self.upload(s, cid, "01_digital_ok.pdf")   # bought by 300112455, not by this client
        inv = self.db.list_invoices(cid)[0]
        self.assertEqual(inv["status"], "review")
        self.assertIn("not_client_invoice", [i["code"] for i in inv["issues"]])


class UnitTests(unittest.TestCase):
    def test_map_document_summary_shape(self):
        fields = map_document({"uuid": "U", "issuerId": "1", "issuerName": "A", "receiverId": "2",
                               "internalId": "N1", "dateTimeIssued": "2026-10-01T09:00:00Z",
                               "netAmount": 100, "total": 114})
        self.assertEqual((fields["invoice_number"], fields["invoice_date"], fields["subtotal"], fields["total"]),
                         ("N1", "2026-10-01", 100.0, 114.0))

    def test_reversed_arabic_repair(self):
        reversed_text = "488-201-736 :يبيرضلا ليجستلا مقر\n5,472.00 :هبيرضلا لماش يلامجالا"
        self.assertTrue(looks_reversed(reversed_text))
        fixed = fix_reversed_arabic(reversed_text)
        self.assertIn("رقم التسجيل الضريبي: 488-201-736", fixed)
        self.assertIn("5,472.00", fixed)
        self.assertFalse(looks_reversed("رقم الفاتورة: 1\nالإجمالي: 100"))

    def test_phone_normalisation(self):
        from app.db import normalize_phone
        self.assertEqual(normalize_phone("01001234567"), "201001234567")
        self.assertEqual(normalize_phone("+20 100 123 4567"), "201001234567")
        self.assertEqual(normalize_phone("٠١٠٠١٢٣٤٥٦٧"), "201001234567")


if __name__ == "__main__":
    unittest.main()
