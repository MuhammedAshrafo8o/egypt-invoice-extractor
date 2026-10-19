"""Excel export of a client's month: VAT summary, purchases, sales, and purchases by supplier.

Every total is an Excel formula over the invoice rows, driven by one explicit column
("محتسبة في الإقرار" = نعم / لا), so the accountant can see — and change — what is counted.
The counting rule matches InvoiceService.summary(): an invoice counts when it is in EGP and
either has no errors or was approved by the accountant.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="134E43")
HEADER_FONT = Font(bold=True, color="FFFFFF", name=FONT)
BODY_FONT = Font(name=FONT)
BOLD = Font(bold=True, name=FONT)
TOTAL_FILL = PatternFill("solid", fgColor="E6EFEC")
EXCLUDED_FILL = PatternFill("solid", fgColor="FFF4D6")
THIN = Side(style="thin", color="D0D7D5")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
DATE_FMT = "dd/mm/yyyy"

MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس",
             "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
STATUS_AR = {"ok": "سليمة", "review": "مراجعة", "error": "خطأ"}
SOURCE_AR = {"upload": "رفع ملف", "whatsapp": "واتساب", "eta_sync": "المنظومة"}

# (header, width, number format)
COLUMNS = [
    ("م", 6, None),
    ("تاريخ الفاتورة", 14, DATE_FMT),
    ("رقم الفاتورة", 18, None),
    ("اسم الطرف الآخر", 34, None),
    ("الرقم الضريبي", 14, "@"),
    ("المبلغ قبل الضريبة", 18, MONEY),
    ("نسبة الضريبة", 12, "0%"),
    ("قيمة الضريبة", 16, MONEY),
    ("الإجمالي", 18, MONEY),
    ("العملة", 9, None),
    ("الحالة", 12, None),
    ("محتسبة في الإقرار", 15, None),
    ("المصدر", 12, None),
    ("ملاحظات", 60, None),
]
COL = {name: i + 1 for i, (name, _, _) in enumerate(COLUMNS)}
L = {name: get_column_letter(i) for name, i in COL.items()}
COUNTED_YES, COUNTED_NO = "نعم", "لا"


def period_label(period: str) -> str:
    try:
        y, m = period.split("-")
        return f"{MONTHS_AR[int(m) - 1]} {y}"
    except (ValueError, IndexError):
        return period


def is_counted(inv: dict) -> bool:
    return (inv["status"] != "error" or bool(inv["approved"])) and (inv.get("currency") or "EGP") == "EGP"


def _status(inv: dict) -> str:
    return "معتمدة" if inv["approved"] else STATUS_AR.get(inv["status"], inv["status"])


def _date(value):
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return value


def _header(ws, columns) -> None:
    for col, (name, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True


def _invoice_sheet(wb: Workbook, title: str, invoices: list[dict], counterparty: str) -> tuple[str, int]:
    """Returns (sheet title, totals row)."""
    ws = wb.create_sheet(title)
    _header(ws, COLUMNS)
    name_key, tax_key = (("supplier_name", "supplier_tax_id") if counterparty == "supplier"
                         else ("buyer_name", "buyer_tax_id"))
    row = 1
    for n, inv in enumerate(sorted(invoices, key=lambda i: (i["invoice_date"] or "", i["id"])), start=1):
        row += 1
        counted = is_counted(inv)
        notes = " | ".join(i.get("message_ar") or i.get("message", "") for i in inv["issues"])
        if (inv.get("currency") or "EGP") != "EGP":
            notes = f"عملة أجنبية ({inv['currency']}) — مش داخلة في المجموع، اتحسبها بالجنيه يدوي. " + notes
        values = [n, _date(inv["invoice_date"]), inv["invoice_number"], inv[name_key], inv[tax_key],
                  inv["subtotal"], inv["vat_rate"], inv["vat_amount"], inv["total"], inv.get("currency") or "EGP",
                  _status(inv), COUNTED_YES if counted else COUNTED_NO,
                  SOURCE_AR.get(inv["channel"], inv["channel"]), notes]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font, cell.border = BODY_FONT, BORDER
            fmt = COLUMNS[col - 1][2]
            if fmt:
                cell.number_format = fmt
            if not counted:
                cell.fill = EXCLUDED_FILL
    last = max(row, 2)
    total_row = row + 1
    if row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{row}"
    ws.cell(row=total_row, column=COL["اسم الطرف الآخر"], value="المجموع (الفواتير المحتسبة بس)")
    flag = L["محتسبة في الإقرار"]
    for name in ("المبلغ قبل الضريبة", "قيمة الضريبة", "الإجمالي"):
        letter = L[name]
        ws.cell(row=total_row, column=COL[name],
                value=f'=SUMIFS({letter}2:{letter}{last},{flag}2:{flag}{last},"{COUNTED_YES}")' if row >= 2 else 0)
    ws.cell(row=total_row, column=COL["م"], value=f'=COUNTIF({flag}2:{flag}{last},"{COUNTED_YES}")' if row >= 2 else 0)
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BORDER
        fmt = COLUMNS[col - 1][2]
        if fmt == MONEY:
            cell.number_format = MONEY
    return title, total_row


def _supplier_sheet(wb: Workbook, purchases: list[dict], p_title: str, p_last: int) -> None:
    """Purchases grouped by supplier — SUMIFS over the purchases sheet, so it follows any edits there."""
    ws = wb.create_sheet("المشتريات حسب المورد")
    cols = [("المورد", 34, None), ("الرقم الضريبي", 14, "@"), ("عدد الفواتير", 12, "0"),
            ("قبل الضريبة", 18, MONEY), ("الضريبة", 16, MONEY), ("الإجمالي", 18, MONEY)]
    _header(ws, cols)
    suppliers: dict = {}
    for inv in purchases:
        if not is_counted(inv):
            continue
        key = inv["supplier_tax_id"] or ""
        entry = suppliers.setdefault(key, {"name": inv["supplier_name"] or "", "vat": 0.0})
        entry["vat"] += inv["vat_amount"] or 0
        if not entry["name"] and inv["supplier_name"]:
            entry["name"] = inv["supplier_name"]
    ref = lambda col: f"'{p_title}'!{L[col]}2:{L[col]}{p_last}"  # noqa: E731
    flag = ref("محتسبة في الإقرار")
    row = 1
    for tax_id, entry in sorted(suppliers.items(), key=lambda kv: -kv[1]["vat"]):
        row += 1
        crit = f'{ref("الرقم الضريبي")},B{row}' if tax_id else f'{ref("اسم الطرف الآخر")},A{row}'
        values = [entry["name"] or "—", tax_id or None,
                  f'=COUNTIFS({crit},{flag},"{COUNTED_YES}")',
                  f'=SUMIFS({ref("المبلغ قبل الضريبة")},{crit},{flag},"{COUNTED_YES}")',
                  f'=SUMIFS({ref("قيمة الضريبة")},{crit},{flag},"{COUNTED_YES}")',
                  f'=SUMIFS({ref("الإجمالي")},{crit},{flag},"{COUNTED_YES}")']
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font, cell.border = BODY_FONT, BORDER
            if cols[col - 1][2]:
                cell.number_format = cols[col - 1][2]
    total = row + 1
    ws.cell(row=total, column=1, value="المجموع")
    for col in (3, 4, 5, 6):
        letter = get_column_letter(col)
        ws.cell(row=total, column=col, value=f"=SUM({letter}2:{letter}{max(row, 2)})" if row >= 2 else 0)
    for col in range(1, len(cols) + 1):
        cell = ws.cell(row=total, column=col)
        cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BORDER
        if cols[col - 1][2] == MONEY:
            cell.number_format = MONEY


def build_workbook(client: dict, period: str, invoices: list[dict], summary: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ملخص الإقرار"
    ws.sheet_view.rightToLeft = True

    purchases = [i for i in invoices if i["direction"] == "purchase"]
    sales = [i for i in invoices if i["direction"] == "sale"]
    p_title, p_total = _invoice_sheet(wb, "المشتريات", purchases, "supplier")
    s_title, s_total = _invoice_sheet(wb, "المبيعات", sales, "buyer")
    _supplier_sheet(wb, purchases, p_title, max(p_total - 1, 2))

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    title = ws.cell(row=1, column=1, value=f"ملخص إقرار ضريبة القيمة المضافة — {period_label(period)}")
    title.font = Font(bold=True, size=14, name=FONT)
    excluded = sum(1 for i in invoices if not is_counted(i))
    rows = [
        (3, "العميل", client["name"], None),
        (4, "الرقم الضريبي", client.get("tax_id") or "—", "@"),
        (5, "الفترة", period_label(period), None),
        (7, "المبيعات قبل الضريبة", f"='{s_title}'!{L['المبلغ قبل الضريبة']}{s_total}", MONEY),
        (8, "ضريبة المخرجات (على المبيعات)", f"='{s_title}'!{L['قيمة الضريبة']}{s_total}", MONEY),
        (9, "المشتريات قبل الضريبة", f"='{p_title}'!{L['المبلغ قبل الضريبة']}{p_total}", MONEY),
        (10, "ضريبة المدخلات (على المشتريات)", f"='{p_title}'!{L['قيمة الضريبة']}{p_total}", MONEY),
        (11, "صافي الضريبة (موجب = مستحق للسداد، سالب = رصيد دائن)", "=B8-B10", MONEY),
        (13, "فواتير الشهر", len(invoices), "0"),
        (14, "فواتير محتسبة في الإقرار", f"='{s_title}'!A{s_total}+'{p_title}'!A{p_total}", "0"),
        (15, "فواتير مش محتسبة (أخطاء أو عملة أجنبية)", excluded, "0"),
        (16, "فواتير محتاجة مراجعة (محتسبة)", summary["review_count"], "0"),
    ]
    for r, label, value, fmt in rows:
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=value)
        a.font, b.font = BOLD, BODY_FONT
        a.border = b.border = BORDER
        if fmt:
            b.number_format = fmt
    for c in ("A11", "B11"):
        ws[c].fill = TOTAL_FILL
        ws[c].font = BOLD
    note = ws.cell(row=18, column=1, value=(
        "الأرقام مساعدة لتجهيز الإقرار، ولازم المحاسب يراجعها قبل التقديم. الفواتير المحتسبة هي اللي "
        "بالجنيه ومافيهاش أخطاء أو اتعتمدت. غيّر عمود \"محتسبة في الإقرار\" في أي شيت والمجاميع هتتحدث."))
    note.font = Font(italic=True, color="666666", name=FONT)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A18:B20")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
